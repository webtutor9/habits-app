import openpyxl, datetime, json, re, sys, unicodedata
src, out = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',s).strip('-')
CODES = {'1':'done','x':'planned','y':'urgent','f':'rest','i':'iso'}
NUMERIC = {'Fastenintervall (h)':'h','Gewicht':'kg','Körperfettanteil (%)':'%','Ruhepuls':'bpm'}
habits, rowmap, section, seen = [], {}, None, set()
for r in range(4, ws.max_row+1):
    if ws.cell(r,1).value: section = ws.cell(r,1).value.strip()
    name = ws.cell(r,2).value
    if not name: continue
    name = name.strip(); hid = slug(name)
    while hid in seen: hid += '-2'
    seen.add(hid)
    h = {'id':hid,'name':name,'section':section}
    if name == 'Delta Vortag': h.update(type='delta', source='gewicht', unit='kg')
    elif name in NUMERIC: h.update(type='number', unit=NUMERIC[name])
    else: h['type'] = 'status'
    habits.append(h); rowmap[r] = h
entries, unknown = {}, {}
base = datetime.date(2026,1,1)
for c in range(3, 368):
    d = (base + datetime.timedelta(c-3)).isoformat()
    for r, h in rowmap.items():
        if h['type'] == 'delta': continue
        v = ws.cell(r,c).value
        if v is None or (isinstance(v,str) and v.strip() in ('','NA')): continue
        if h['type'] == 'number':
            try: val = round(float(v), 2)
            except: unknown[str(v)] = unknown.get(str(v),0)+1; continue
        else:
            k = str(v).strip().lower()
            if isinstance(v,(int,float)) and v == 1: k = '1'
            if k not in CODES: unknown[k] = unknown.get(k,0)+1; continue
            val = CODES[k]
        entries.setdefault(d, {})[h['id']] = val
data = {'version':1,'habits':habits,'entries':entries,
        'updatedAt':datetime.datetime.now(datetime.timezone.utc).isoformat()}
json.dump(data, open(out,'w'), ensure_ascii=False, indent=1)
print(len(habits),'habits', len(entries),'days', 'unknown:',unknown)
